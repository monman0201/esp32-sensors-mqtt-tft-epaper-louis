# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"D:\monman\esp32\ESP32實習")
OUT = ROOT / "0914物聯網成果報告_最終整合版.docx"
DG = ROOT / "_report_diagrams"
PH = ROOT / "成果照片"
FONT = "Microsoft JhengHei"
NAVY, TEAL, GOLD, GRAY = "17324D", "168AAD", "E9A23B", "5B6770"

def font(run, size=11, color="222222", bold=False):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    for k in ("eastAsia", "ascii", "hAnsi"):
        rpr.rFonts.set(qn("w:" + k), FONT)
    run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color); run.bold = bold

def para(doc, text="", size=11, color="222222", bold=False, align=None, after=6):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.2
    if align is not None: p.alignment = align
    font(p.add_run(text), size, color, bold); return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    font(p.add_run(text), 16 if level == 1 else 12.5, TEAL if level == 1 else NAVY, True); return p

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(3); font(p.add_run(text), 10.5); return p

def table(doc, rows, widths):
    t = doc.add_table(rows=len(rows), cols=len(widths)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            c = t.cell(i, j); c.width = Inches(widths[j]); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), NAVY); c._tc.get_or_add_tcPr().append(shd)
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2); font(p.add_run(str(value)), 9.2, "FFFFFF" if i == 0 else "222222", i == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2); return t

def picture(doc, path, caption, width=5.8):
    if path.exists():
        doc.add_picture(str(path), width=Inches(width)); para(doc, caption, 9, GRAY, False, WD_ALIGN_PARAGRAPH.CENTER, 8)

doc = Document(); sec = doc.sections[0]
sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.68); sec.left_margin = Inches(.82); sec.right_margin = Inches(.82)
styles = doc.styles; styles["Normal"].font.name = FONT; styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT); styles["Normal"].font.size = Pt(11)
for name, size, color in (("Heading 1",16,TEAL),("Heading 2",12.5,NAVY),("Heading 3",11.5,GRAY)):
    s=styles[name]; s.font.name=FONT; s._element.rPr.rFonts.set(qn("w:eastAsia"),FONT); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(header.add_run("ESP32 物聯網環境監測與雲端控制系統｜期末成果報告"),9,GRAY)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(footer.add_run("0914物聯網成果報告｜2026 年 9 月 14 日"),9,GRAY)

# Cover
para(doc,"物聯網環境監測與雲端控制系統",28,NAVY,True,WD_ALIGN_PARAGRAPH.CENTER,10)
para(doc,"ESP32 × 感測器 × OLED／ILI9225 TFT × ThingSpeak／Google Sheets／LINE／MQTT",13,TEAL,False,WD_ALIGN_PARAGRAPH.CENTER,28)
para(doc,"期末學習成果報告",21,GOLD,True,WD_ALIGN_PARAGRAPH.CENTER,45)
para(doc,"專案整合範圍：ESP32 物聯網實作\n報告日期：2026 年 9 月 14 日\n使用語言：Arduino C++／繁體中文（UTF-8）",11,GRAY,False,WD_ALIGN_PARAGRAPH.CENTER,18)
para(doc,"本專案以 ESP32 為核心，整合環境感測、現場顯示、雲端資料紀錄、異常通知、MQTT 雙向通訊與輸出控制，完成一套可觀測、可分析、可互動的物聯網環境監測系統。",12,"333333",False,WD_ALIGN_PARAGRAPH.CENTER,0)
doc.add_page_break()

heading(doc,"摘要")
para(doc,"本報告完整整理 ESP32 物聯網實作成果。系統以 DHT11 量測溫度與濕度，以光敏電阻取得環境亮度；資料可在 OLED 與 ILI9225 TFT 上即時顯示，也能透過 ThingSpeak 與 Google Sheets 建立雲端紀錄，並以 LINE 發送異常通知。最新整合版本加入 MQTT 感測資料發布、NTP 台灣時間、四頁 TFT 儀表板、歷史趨勢圖與三路 LED 遠端控制，使系統具備資料分析與互動控制能力。")
table(doc, [["面向","整合成果"],["感測","DHT11 溫度／濕度、光敏電阻亮度"],["現場顯示","128×64 OLED 與 ILI9225 TFT 四頁儀表板"],["雲端資料","ThingSpeak、Google Sheets、MQTT Broker"],["通知與控制","LINE 異常通知、MQTT 三色 LED 控制"],["時間與操作","NTP UTC+8、GPIO0 四頁換頁"]],[1.7,5.1])
heading(doc,"一、學習目標")
para(doc,"這次 ESP32 實習的學習目標，是從基礎的感測器讀取開始，逐步完成一套可以實際運作的物聯網系統。我希望透過實作了解感測資料如何被讀取、顯示、上傳到雲端，並進一步完成通知與設備控制。")
for x in ["學會使用 DHT11 與光敏電阻讀取溫度、濕度及亮度。","學會使用 OLED 與 ILI9225 TFT 設計現場顯示介面。","了解 ThingSpeak、Google Sheets 與 MQTT 的資料傳輸方式。","實作 LINE 異常通知與 MQTT 三色 LED 控制。","練習 Wi-Fi、NTP、JSON、網路斷線重連與按鍵換頁等整合功能。"]: bullet(doc,x)
doc.add_page_break()

heading(doc,"二、整合系統架構")
para(doc,"系統以 ESP32 作為資料處理與通訊核心。感測層取得環境資料，顯示層提供現場回饋，雲端層負責歷史紀錄與即時訊息，應用層完成通知與輸出控制。ThingSpeak、Google Sheets、LINE 與 MQTT 都以同一組感測資料為基礎，對應資料分析、整理、通知與控制需求。")
picture(doc,DG/"architecture_final_integrated.png","圖 1　ESP32 物聯網環境監測與雲端控制完整系統架構圖（imagegen 製作）",6.35)
heading(doc,"三、硬體架構與接線")
table(doc,[["元件","GPIO／介面","功能"],["ESP32","主控制器","感測、資料處理、網路通訊與輸出控制"],["DHT11","GPIO14","量測溫度與濕度"],["光敏電阻","GPIO33（ADC）","量測類比亮度並換算百分比"],["OLED","SDA GPIO21、SCL GPIO22","顯示感測數值與傳送狀態"],["ILI9225 TFT","RST26、RS25、CS16、MOSI23、SCK18","四頁儀表板、Gauge 與歷史折線圖"],["三色 LED","綠 GPIO15、黃 GPIO2、紅 GPIO4","依 MQTT 指令分別開關"],["換頁按鍵","GPIO0，INPUT_PULLUP","循環切換 P1/4～P4/4"]],[1.55,2.55,2.7])
para(doc,"硬體由麵包板完成原型接線，DHT11 與光敏電阻負責資料輸入，OLED／TFT 負責顯示，三顆 LED 作為 MQTT 控制結果的實體回饋。GPIO0 使用按鍵中斷並搭配防彈跳，避免短按事件被網路處理流程遺漏。")
picture(doc,PH/"20260810_091121.jpg","圖 3　ESP32、OLED、DHT11 與光敏電阻實體接線成果",5.35)
doc.add_page_break()

heading(doc,"四、軟體流程與資料傳遞")
para(doc,"系統開機後初始化 GPIO、感測器、顯示器、SPI／I²C 與網路服務；主迴圈持續讀取感測值並更新顯示，再依服務週期送出資料。網路中斷時重新建立 Wi-Fi 或 MQTT 連線；收到控制訊息時立即更新對應 LED。")
picture(doc,DG/"flow_final_integrated.png","圖 4　ESP32 物聯網感測、雲端服務與控制完整流程圖（imagegen 製作）",6.35)
table(doc,[["流程環節","整合行為"],["感測","DHT11 讀取溫度／濕度，GPIO33 讀取亮度"],["現場顯示","OLED 即時資訊；TFT 首頁與三種趨勢頁"],["雲端資料","ThingSpeak 15 秒、Google Sheets 10 秒、MQTT 10 秒"],["通知／控制","LINE 依門檻通知；MQTT JSON 控制三色 LED"],["資料分析","TFT 保存最多 60 筆資料，繪製最近 10 分鐘折線圖"]],[1.7,5.1])
doc.add_page_break()

heading(doc,"五、功能實作與整合")
heading(doc,"5.1 感測與 OLED 即時顯示",2)
para(doc,"OLED 以三個資訊區塊呈現溫度、濕度與亮度；中文欄位名稱與單位清楚顯示，右側以 UP／OK 或 TX／OK 表示資料傳送狀態，方便確認感測與網路服務是否正常。")
picture(doc,PH/"螢幕擷取畫面 2026-08-10 094915.png","圖 6　感測資料與 OLED 顯示成果",4.5)
heading(doc,"5.2 ThingSpeak 雲端資料上傳",2)
para(doc,"ESP32 透過 Wi-Fi 以 HTTP GET 將感測資料送至 ThingSpeak Channel，建立瀏覽器可查看的歷史圖表。上傳週期為 15 秒，field1 為溫度、field2 為濕度、field3 為亮度；測試取得 HTTP code 200，並建立 entry id 7。此功能讓現場即時顯示延伸至雲端歷史分析。")
table(doc,[["欄位","內容","用途"],["field1","溫度（°C）","溫度歷史曲線"],["field2","濕度（%）","濕度歷史曲線"],["field3","亮度（%）","亮度歷史曲線"]],[1.4,2.2,3.2])
picture(doc,PH/"thingspeak上傳成果(305教室).png","圖 7　ThingSpeak 雲端資料上傳與歷程圖表成果",5.65)
heading(doc,"5.3 Google Sheets 資料紀錄",2)
para(doc,"Google Sheets 版本使用 HTTPS 連線至 Google Apps Script Web App，以 URL 參數傳遞 CSV 格式資料，欄位包含時間戳記、溫度、濕度與亮度，每 10 秒傳送一筆，並在 OLED 顯示 UP／OK。此功能適合資料整理、匯出與後續統計。")
picture(doc,PH/"螢幕擷取畫面 2026-08-10 113831.png","圖 8　Google Sheets 感測資料紀錄成果",4.4)
picture(doc,PH/"螢幕擷取畫面 2026-08-10 113849.png","圖 9　雲端資料圖表呈現成果",5.6)
heading(doc,"5.4 LINE 異常通知",2)
para(doc,"LINE 版本使用 WiFiClientSecure 連線至 LINE Messaging API；當溫度超過 28°C 或濕度超過 70% 時推送異常訊息，並以 30 秒為通知間隔，避免同一異常狀態重複通知。")
picture(doc,PH/"Screenshot_20260810_160929_LINE.jpg","圖 10　LINE 異常通知實際接收畫面",3.7)
heading(doc,"5.5 MQTT 感測資料與三色 LED 控制",2)
para(doc,"整合版本使用 mqttgo.io、Port 1883 與 MQTT 3.1.1。感測資料主題為 louis/class305/data，每 10 秒發布 JSON：{\"temp\":25,\"humi\":65,\"light\":95}。綠、黃、紅 LED 分別訂閱 louis/class305/ctrl/gled、louis/class305/ctrl/yled、louis/class305/ctrl/rled，並使用 ArduinoJson 解析 {\"gled\":\"on\"} 等控制訊息。")
picture(doc,PH/"20260831_153358.jpg","圖 11　環境監測裝置、感測器與顯示器整合測試成果",5.7)
picture(doc,PH/"20260831_154454.jpg","圖 12　麵包板接線與環境資料顯示測試成果",5.7)
picture(doc,PH/"20260831_154459.jpg","圖 13　感測資料與輸出元件整合測試成果",5.7)
picture(doc,PH/"20260914_115835.jpg","圖 14　MQTTGO Broker 接收感測資料成果",5.65)
heading(doc,"5.6 ILI9225 TFT 四頁儀表板與 NTP",2)
para(doc,"ILI9225 TFT 以 176×220 顯示四頁介面：P1/4 為環境摘要與 Wi-Fi／MQTT 狀態；P2/4 為溫度 Gauge 與最近 10 分鐘折線圖；P3/4 為濕度 Gauge 與折線圖；P4/4 為亮度 Gauge 與折線圖。NTP 以 UTC+8 顯示台灣日期時間，GPIO0 按鍵可循環切換頁面。")
picture(doc,PH/"20260914_134740.jpg","圖 15　ESP32 環境監測裝置與 TFT 顯示運作成果",5.7)
picture(doc,PH/"20260914_135704.jpg","圖 16　ILI9225 TFT、感測器與三色 LED 整合成果",5.7)
doc.add_page_break()

heading(doc,"六、學習成果")
table(doc,[["測試項目","測試條件／觀察","結果"],["感測讀值","DHT11 與光敏電阻持續讀取","通過"],["OLED 顯示","中文欄位、數值、單位與狀態","通過"],["ThingSpeak","15 秒上傳 field1～field3，HTTP 200","通過"],["Google Sheets","10 秒寫入溫度、濕度、亮度","通過／依網路狀態而定"],["LINE 通知","溫度 > 28°C 或濕度 > 70%","通過；每 30 秒一次"],["MQTT","感測資料發布與三路 LED 訂閱控制","通過"],["TFT 四頁","GPIO0 換頁、Gauge、10 分鐘趨勢圖","通過"],["網路中斷","Wi-Fi／MQTT 斷線後重新連線","已納入程式邏輯"]],[2.0,3.3,1.7])
para(doc,"透過這次實作，我完成了從感測資料讀取、OLED／TFT 顯示、ThingSpeak 與 Google Sheets 雲端紀錄，到 LINE 通知及 MQTT 遠端控制的整合。這些成果讓我更清楚理解不同物聯網服務的用途，也實際練習了網路連線、JSON 資料格式、API、MQTT 與嵌入式顯示介面的整合。")
heading(doc,"七、學習心得")
para(doc,"這次實習讓我體會到，物聯網不只是把感測器接上 ESP32，而是要讓硬體、程式、網路與使用者介面一起協調運作。剛開始我比較熟悉的是單一感測器與 OLED 顯示，後來逐步加入 ThingSpeak、Google Sheets、LINE、MQTT、TFT 和 LED 控制後，才真正了解一個完整系統需要處理許多細節，例如資料格式、傳送週期、網路中斷、畫面配置與硬體腳位。")
para(doc,"在製作過程中，我學到如何閱讀程式執行結果、檢查接線與追蹤網路訊息，也學會把一個大功能拆成感測、顯示、上傳、通知和控制等部分再逐步整合。最後完成四頁 TFT 儀表板與 MQTT 控制時，讓我很有成就感。這次經驗增加了我對嵌入式系統與物聯網應用的信心，也讓我知道未來學習更大型專題時，應該先規劃系統架構，再逐項完成與驗證。")
heading(doc,"附錄：專案檔案與資料欄位")
table(doc,[["項目","內容"],["主要程式","25_dht_line、32_ili_mqtt_ctrl_page"],["感測資料欄位","時間戳記、溫度、濕度、亮度"],["ThingSpeak 欄位","field1 溫度、field2 濕度、field3 亮度"],["MQTT 感測主題","louis/class305/data"],["MQTT 控制主題","louis/class305/ctrl/gled、yled、rled"],["重要門檻","溫度 > 28°C 或濕度 > 70%；LINE 通知間隔 30 秒"]],[2.2,4.8])

doc.core_properties.title="0914物聯網成果報告"; doc.core_properties.subject="ESP32 物聯網環境監測與雲端控制系統期末學習成果"; doc.core_properties.comments="UTF-8 製作，整合原有圖文與 0914 最新成果。"
doc.save(str(OUT)); print(OUT)
