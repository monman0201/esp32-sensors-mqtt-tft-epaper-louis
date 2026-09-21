# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"D:\monman\esp32\ESP32實習")
SOURCE = ROOT / "0810物聯網成果報告.docx"
OUTPUT = ROOT / "0914物聯網成果報告_整合版.docx"
DIAGRAMS = ROOT / "_report_diagrams"
PHOTOS = ROOT / "成果照片"
FONT = "Microsoft JhengHei"
NAVY, TEAL, GRAY = "17324D", "168AAD", "5B6770"


def fmt(run, size=11, color="222222", bold=False):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    for key in ("eastAsia", "ascii", "hAnsi"):
        rpr.rFonts.set(qn("w:" + key), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def para(text="", size=11, color="222222", bold=False, align=None):
    d = Document()
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.2
    if align is not None:
        p.alignment = align
    fmt(p.add_run(text), size, color, bold)
    return p


def heading(text, level=2):
    d = Document()
    p = d.add_paragraph(style=f"Heading {level}")
    fmt(p.add_run(text), 12.5 if level == 2 else 11.5, NAVY, True)
    return p


def bullet(text):
    d = Document()
    p = d.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    fmt(p.add_run(text), 10.5)
    return p


TARGET_DOC = None


def picture(path, caption, width=6.2):
    # Create the picture in the target document first, so its relationship is
    # registered in the target .docx package before the paragraph is moved.
    TARGET_DOC.add_picture(str(path), width=Inches(width))
    pic_p = TARGET_DOC.paragraphs[-1]
    p = TARGET_DOC.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt(p.add_run(caption), 9, GRAY)
    return [pic_p._p, p._p]


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def table(rows, widths):
    d = Document()
    t = d.add_table(rows=len(rows), cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            c = t.cell(i, j)
            c.width = Inches(widths[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                shade(c, NAVY)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            fmt(p.add_run(str(value)), 9.2, "FFFFFF" if i == 0 else "222222", i == 0)
    return [t._tbl]


def insert_after(anchor, elements):
    current = anchor._p
    for element in elements:
        if hasattr(element, "_p"):
            element = element._p
        current.addnext(element)
        current = element


def insert_before(anchor, elements):
    current = anchor._p
    for element in reversed(elements):
        if hasattr(element, "_p"):
            element = element._p
        current.addprevious(element)
        current = element


def find_para(doc, text):
    for p in doc.paragraphs:
        if text in p.text:
            return p
    raise ValueError(f"anchor not found: {text}")


doc = Document(str(SOURCE))
TARGET_DOC = doc

# Update report date and visible report naming while retaining historical photo dates.
for p in list(doc.paragraphs) + [p for s in doc.sections for p in (s.header.paragraphs + s.footer.paragraphs)]:
    for run in p.runs:
        run.text = run.text.replace("2026 年 8 月 10 日", "2026 年 9 月 14 日")
        run.text = run.text.replace("0810物聯網成果報告", "0914物聯網成果報告")
        run.text = run.text.replace("結案報告", "成果整合報告")

# Integrate the current architecture into the existing architecture section.
arch = find_para(doc, "二、系統架構")
insert_after(arch, [
    para("專案目前已由早期 OLED／Google Sheets／LINE 版本，延伸為 ILI9225 TFT 四頁環境監測儀表板。DHT11 與光敏電阻輸入 ESP32，ESP32 透過 Wi-Fi 連線 MQTT Broker 與 NTP 時間服務；本地端輸出至 TFT，並以 MQTT 指令控制三色 LED。"),
    *picture(DIAGRAMS / "architecture_0914_imagegen.png", "圖 9　整合後 ESP32 物聯網系統架構圖（imagegen 製作）"),
    *table([
        ["元件／服務", "目前整合功能"],
        ["DHT11、光敏電阻", "溫度、濕度與亮度感測"],
        ["ESP32 + ILI9225 TFT", "四頁首頁、Gauge 與最近 10 分鐘折線圖"],
        ["MQTT mqttgo.io", "每 10 秒發布 JSON 感測資料，並接收 LED 控制指令"],
        ["NTP、GPIO0", "UTC+8 日期時間與四頁換頁按鍵"],
    ], [2.2, 4.5]),
])

# Integrate the current flow into the existing software-flow section.
flow = find_para(doc, "資料流程說明")
insert_after(flow, [
    *picture(DIAGRAMS / "flow_0914_imagegen.png", "圖 10　整合後 ESP32 程式運作流程圖（imagegen 製作）"),
    para("最新版本 32_ili_mqtt_ctrl_page 的流程是：初始化 GPIO、SPI TFT、DHT11、ADC、Wi-Fi 與 NTP；連線 MQTT 後持續讀取感測值、更新目前頁面、每 10 秒發布 JSON，並將最多 60 筆資料用於 10 分鐘趨勢圖。收到三個 LED MQTT 主題的 JSON 指令時，callback 立即更新對應 GPIO；Wi-Fi 或 MQTT 中斷時則自動重連。"),
    bullet("P1/4 顯示環境摘要與 Wi-Fi／MQTT／台灣時間；P2/4～P4/4 分別顯示溫度、濕度、亮度 Gauge 與趨勢圖。"),
    bullet("GPIO0 採 FALLING 中斷搭配防彈跳，按一次循環切換四個頁面。"),
])

# Make the ThingSpeak learning stage explicit. The original report already
# contains its screenshot; this text connects it to the documented project
# evolution instead of leaving it as an unlabeled cloud-upload image.
thingspeak_anchor = find_para(doc, "5.2 Google Sheets 上傳")
insert_before(thingspeak_anchor, [
    heading("ThingSpeak 雲端資料上傳學習成果（學習歷程）", 2),
    para("在 Google Sheets 與 LINE 之前，專案先完成 ThingSpeak 雲端資料上傳。ESP32 透過 Wi-Fi 以 HTTP GET 將環境資料送至 ThingSpeak Channel，建立可長期查看的雲端歷史圖表，讓溫度、濕度與亮度不只在 OLED 上即時顯示，也能在瀏覽器中追蹤變化。"),
    *table([
        ["ThingSpeak 欄位", "資料內容", "用途"],
        ["field1", "溫度（°C）", "觀察環境溫度趨勢"],
        ["field2", "濕度（%）", "觀察環境濕度趨勢"],
        ["field3", "亮度（%）", "觀察光線變化"],
    ], [1.5, 2.0, 3.2]),
    para("ThingSpeak 版本的上傳週期為 15 秒；測試曾取得 HTTP code 200，並成功建立資料 entry，證明感測資料已由 ESP32 經 Wi-Fi 傳送至雲端。原報告中的「雲端資料上傳與歷程呈現成果」照片即為此階段的實作佐證。"),
    para("此階段建立了後續 Google Sheets、LINE 與 MQTT 功能的基礎：ThingSpeak 偏重歷史資料視覺化，Google Sheets 便於資料整理，LINE 著重異常通知，而 0914 版本進一步加入 MQTT 即時發布與雙向控制。"),
])

# Add the latest physical and MQTT evidence inside the existing results section.
test = find_para(doc, "六、測試與成果")
insert_after(test, [
    para("在保留舊版 OLED、雲端紀錄與 LINE 通知成果的基礎上，0914 版本進一步完成 TFT 儀表板與 MQTT 雙向控制，形成由感測、顯示、雲端傳輸到輸出控制的整合成果。"),
    *table([
        ["0914 測試項目", "結果"],
        ["ESP32 Dev Module 編譯", "通過；Flash 約 59%、RAM 約 14%"],
        ["Wi-Fi／NTP", "可連線並以 UTC+8 顯示日期時間"],
        ["MQTT 感測資料", "每 10 秒發布至 louis/class305/data"],
        ["四頁 TFT 與歷史圖", "GPIO0 可切換 P1/4～P4/4，最多記錄 60 筆"],
        ["三色 LED 控制", "可由三個 MQTT 主題分別控制 GPIO15、GPIO2、GPIO4"],
    ], [2.3, 4.4]),
    *picture(PHOTOS / "20260914_135704.jpg", "圖 11　ESP32、ILI9225 TFT、感測器與三色 LED 實體整合成果", 5.7),
    *picture(PHOTOS / "20260914_115835.jpg", "圖 12　MQTTGO Broker 接收 louis/class305/data 感測資料成果", 5.7),
    *picture(PHOTOS / "20260914_134740.jpg", "圖 13　ESP32 環境監測裝置運作與 TFT 顯示成果", 5.7),
])

# Update the conclusion so it describes the combined old and new results.
for p in doc.paragraphs:
    if "本專案以 ESP32 完成環境監測物聯網原型" in p.text:
        for run in p.runs:
            run.text = run.text.replace(
                "整合 DHT11、光敏電阻、OLED、Google Sheets 與 LINE 通知。",
                "整合 DHT11、光敏電阻、OLED、Google Sheets、LINE、ILI9225 TFT 與 MQTT。"
            )
            run.text = run.text.replace(
                "系統具備即時顯示、週期上傳與異常推播三項核心能力",
                "系統具備即時顯示、週期上傳、異常推播、MQTT 雙向控制與歷史趨勢分析等能力"
            )

doc.core_properties.title = "0914物聯網成果報告"
doc.core_properties.subject = "ESP32 物聯網成果整合報告"
doc.core_properties.comments = "以 UTF-8 編碼程式與繁體中文內容整合 0810 舊成果及 0914 最新進度。"
doc.save(str(OUTPUT))
print(OUTPUT)
