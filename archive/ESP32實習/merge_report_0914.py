# -*- coding: utf-8 -*-
"""Merge the 2026-09-14 ESP32 progress into the original 0810 report."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"D:\monman\esp32\ESP32實習")
SOURCE = ROOT / "0810物聯網成果報告.docx"
OUTPUT = ROOT / "0914物聯網成果報告.docx"
DIAGRAMS = ROOT / "_report_diagrams"
PHOTOS = ROOT / "成果照片"

FONT = "Microsoft JhengHei"
NAVY = "17324D"
TEAL = "168AAD"
GRAY = "5B6770"


def set_font(run, size=11, color="222222", bold=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def add_para(doc, text="", size=11, color="222222", bold=False, align=None, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.2
    if align is not None:
        p.alignment = align
    set_font(p.add_run(text), size, color, bold)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_font(p.add_run(text), 16 if level == 1 else 12.5, TEAL if level == 1 else NAVY, True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(text), 10.5)
    return p


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_table(doc, rows, widths):
    table = doc.add_table(rows=len(rows), cols=len(widths))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                shade(cell, NAVY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(value))
            set_font(r, 9.5, "FFFFFF" if i == 0 else "222222", i == 0)
    return table


def add_picture(doc, path, caption, width=6.2):
    if path.exists():
        doc.add_picture(str(path), width=Inches(width))
        add_para(doc, caption, 9, GRAY, False, WD_ALIGN_PARAGRAPH.CENTER, 8)


def update_header_footer(doc):
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            if paragraph.text:
                paragraph.text = ""
                set_font(paragraph.add_run("ESP32 物聯網成果報告｜0914 新進度整合"), 9, GRAY)
        for paragraph in section.footer.paragraphs:
            if paragraph.text:
                paragraph.text = ""
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_font(paragraph.add_run("0914物聯網成果報告"), 9, GRAY)


doc = Document(str(SOURCE))
update_header_footer(doc)
doc.core_properties.title = "0914物聯網成果報告"
doc.core_properties.subject = "ESP32 物聯網環境監測與 MQTT 控制系統：舊成果與最新進度整合"
doc.core_properties.comments = "本文件以 UTF-8 編碼原始程式與繁體中文內容製作。"

doc.add_page_break()
add_heading(doc, "九、2026 年 9 月 14 日新進度整合", 1)
add_para(doc, "本章承接前述 0810 成果，保留原有的 DHT11、光敏電阻、OLED、Google Sheets 與 LINE 實作；本次再加入 ILI9225 TFT 儀表板、MQTT 雲端通訊、NTP 時間同步、四頁趨勢圖、GPIO0 換頁與三路 LED 控制。最新接續版本為 32_ili_mqtt_ctrl_page，程式已完成 ESP32 Dev Module 編譯驗證。")
add_table(doc, [
    ["項目", "0914 最新成果"],
    ["目前版本", "32_ili_mqtt_ctrl_page；ILI9225 TFT 四頁版"],
    ["感測資料", "DHT11 溫度／濕度；GPIO33 光敏電阻亮度"],
    ["網路服務", "Wi-Fi、MQTT mqttgo.io:1883、NTP UTC+8"],
    ["MQTT 資料", "louis/class305/data；每 10 秒發布 {temp, humi, light}"],
    ["控制功能", "三個 MQTT 主題分別控制綠、黃、紅 LED"],
    ["顯示功能", "首頁、溫度、濕度、亮度四頁；Gauge 與最近 10 分鐘折線圖"],
], [1.6, 5.2])

add_heading(doc, "9.1 最新系統架構圖", 2)
add_picture(doc, DIAGRAMS / "architecture_0914_imagegen.png", "圖 9　0914 最新 ESP32 物聯網系統架構圖（imagegen 製作）")
add_para(doc, "最新架構以 ESP32 為核心：左側接收 DHT11 與光敏電阻資料，右側透過 Wi-Fi 連線到 MQTT Broker 與 NTP 時間服務；本地端則將資料呈現於 ILI9225 TFT，並依 MQTT 指令驅動三色 LED。GPIO0 按鍵用於切換四個顯示頁面。")

add_heading(doc, "9.2 硬體接線與功能對照", 2)
add_table(doc, [
    ["元件", "GPIO／介面", "功能"],
    ["ILI9225 TFT", "RST 26、RS 25、CS 16、MOSI 23、SCK 18", "顯示四頁儀表板與趨勢圖"],
    ["DHT11", "GPIO14", "讀取溫度與濕度"],
    ["光敏電阻", "GPIO33 ADC", "讀取亮度並換算 0～100%"],
    ["綠／黃／紅 LED", "GPIO15／2／4", "依 MQTT 指令分別開關"],
    ["換頁按鍵", "GPIO0，INPUT_PULLUP", "循環切換 P1/4～P4/4"],
], [1.5, 2.6, 2.7])

add_heading(doc, "9.3 最新程式流程圖", 2)
add_picture(doc, DIAGRAMS / "flow_0914_imagegen.png", "圖 10　0914 最新程式運作流程圖（imagegen 製作）")
add_para(doc, "程式開機後初始化 SPI、TFT、感測器、按鍵與 LED，接著建立 Wi-Fi、NTP 與 MQTT 連線。主迴圈持續讀取感測值，更新目前頁面、每 10 秒發布 JSON，並將資料寫入最多 60 筆的歷史陣列；收到 LED 的 MQTT JSON 指令時，callback 會立即更新對應 GPIO。若網路中斷，程式會自動重新連線。")

add_heading(doc, "9.4 ILI9225 四頁介面", 2)
for text in [
    "P1/4 環境監控首頁：顯示溫度、濕度、亮度資訊卡、Wi-Fi／MQTT 狀態與台灣時間；濕度超過 65% 時以閃爍方式警示。",
    "P2/4 溫度：10～40°C 半圓 Gauge，加上最近 10 分鐘、最多 60 筆的溫度折線圖。",
    "P3/4 濕度：0～100% Gauge，並以黃／綠／紅色區分偏乾、舒適與偏濕區間。",
    "P4/4 亮度：0～100% Gauge，並以藍／黃／紅色區分偏暗、中等與偏亮區間。",
    "GPIO0 使用 FALLING 中斷偵測短按，再以防彈跳機制循環切換四頁。",
]:
    add_bullet(doc, text)

add_heading(doc, "9.5 0914 實體與 MQTT 成果照片", 2)
add_picture(doc, PHOTOS / "20260914_135704.jpg", "圖 11　ESP32、ILI9225 TFT、DHT11、光敏電阻與三色 LED 實體整合成果", 5.7)
add_picture(doc, PHOTOS / "20260914_115835.jpg", "圖 12　MQTTGO Broker 接收 louis/class305/data 感測資料成果", 5.7)
add_picture(doc, PHOTOS / "20260914_134740.jpg", "圖 13　ESP32 環境監測裝置運作與 TFT 顯示成果", 5.7)

add_heading(doc, "9.6 版本演進與整合成果", 2)
add_table(doc, [
    ["階段", "主要功能", "本次整合狀態"],
    ["22～25", "DHT11／光敏電阻、OLED、ThingSpeak／Google Sheets／LINE", "保留於本報告前段，作為早期成果"],
    ["27", "MQTT 感測資料與 LED 控制", "延伸為三路獨立控制主題"],
    ["28～31", "ILI9225 TFT、MQTT、ArduinoJson", "作為目前四頁版基礎"],
    ["32", "NTP、GPIO0 換頁、四頁 Gauge 與歷史折線圖", "0914 最新完成版本"],
], [1.0, 3.0, 3.2])
add_para(doc, "整合後的專案形成由『現場顯示與雲端紀錄』逐步發展為『TFT 儀表板、MQTT 雙向控制與歷史趨勢分析』的完整學習成果。舊版功能仍可作為實作歷程，新版則代表目前可展示與持續擴充的主要版本。")

add_heading(doc, "十、最新測試結果與後續建議", 1)
add_table(doc, [
    ["測試項目", "結果"],
    ["ESP32 Dev Module 編譯", "通過；Flash 約 59%、RAM 約 14%"],
    ["Wi-Fi／NTP", "可連線並以 UTC+8 顯示日期時間"],
    ["MQTT 感測資料", "每 10 秒發布至 louis/class305/data"],
    ["TFT 四頁與趨勢圖", "頁面可由 GPIO0 循環切換，歷史資料最多 60 筆"],
    ["三色 LED 控制", "可透過三個 MQTT 主題分別控制 GPIO15、GPIO2、GPIO4"],
    ["後續建議", "重新燒錄最新程式，等待約 10 分鐘確認三種折線圖逐步填滿"],
], [2.3, 4.9])
add_para(doc, "本報告以原 0810 成果為底稿追加 0914 新進度，圖片與文字皆依專案資料夾內的 UTF-8 程式、交接紀錄及成果照片整理。")

doc.save(str(OUTPUT))
print(OUTPUT)
